#!/usr/bin/env python3
"""
p-BART Minimum Draw Threshold & Convergence Exploration Script.
Evaluates multiple MCMC draw/tune configurations across Simple 2, Complex 2, and Complex 3
to identify the minimum draw threshold for statistical performance (RMSE, coverage),
MCMC convergence (R-hat), effective sample size (ESS), and runtime efficiency.
Outputs results to results/min_draw_threshold_study/.
"""

import os
os.environ["OMP_NUM_THREADS"] = "1"
os.environ["OPENBLAS_NUM_THREADS"] = "1"
os.environ["VECLIB_MAXIMUM_THREADS"] = "1"
os.environ["NUMBA_NUM_THREADS"] = "1"

import sys
import time
import json
from pathlib import Path
from typing import Dict, Any, List
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Ensure src in sys.path
SRC_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SRC_DIR))

from config.scenarios import SCENARIOS
from validation.dgp import generate_scenario_data
from validation.metrics import compute_interval_metrics
from validation.runners.pbart_runner import run_pbart
from validation.runners.rbart_runner import run_rbart

BASE_DIR = SRC_DIR.parent
OUTPUT_DIR = BASE_DIR / "results" / "min_draw_threshold_study"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# Parameter grids to test
CONFIGURATIONS = {
    "cfg1_draws250_tune100": {
        "trees": 25,
        "draws": 250,
        "tune": 100,
        "chains": 4,
        "thin": 2,
        "thin_idata": False,
        "update_fraction": 1.0,
        "split_method": "empirical",
        "label": "25 trees | 250 draws (thin 2 -> 500 draws) | 100 tune"
    },
    "cfg2_draws500_tune250": {
        "trees": 25,
        "draws": 500,
        "tune": 250,
        "chains": 4,
        "thin": 5,
        "thin_idata": False,
        "update_fraction": 1.0,
        "split_method": "empirical",
        "label": "25 trees | 500 draws (thin 5 -> 400 draws) | 250 tune"
    },
    "cfg3_draws1000_tune250_thin5": {
        "trees": 50,
        "draws": 1000,
        "tune": 250,
        "chains": 4,
        "thin": 5,
        "thin_idata": False,
        "update_fraction": 1.0,
        "split_method": "empirical",
        "label": "50 trees | 1,000 draws (thin 5 -> 800 draws) | 250 tune"
    },
    "cfg3_draws1000_tune250_thin4": {
        "trees": 50,
        "draws": 1000,
        "tune": 250,
        "chains": 4,
        "thin": 4,
        "thin_idata": False,
        "update_fraction": 1.0,
        "split_method": "empirical",
        "label": "50 trees | 1,000 draws (thin 4 -> 1,000 draws) | 250 tune"
    },
    "cfg4_draws2500_tune250": {
        "trees": 50,
        "draws": 2500,
        "tune": 250,
        "chains": 4,
        "thin": 10,
        "thin_idata": False,
        "update_fraction": 1.0,
        "split_method": "empirical",
        "label": "50 trees | 2,500 draws (thin 10 -> 1,000 draws) | 250 tune (Prod)"
    },
    "cfg5_draws4000_tune500": {
        "trees": 50,
        "draws": 4000,
        "tune": 500,
        "chains": 4,
        "thin": 10,
        "thin_idata": False,
        "update_fraction": 1.0,
        "split_method": "empirical",
        "label": "50 trees | 4,000 draws (thin 10 -> 1,600 draws) | 500 tune"
    }
}

TARGET_SCENARIOS = ["simple_2_1", "complex_2", "complex_3"]
SAMPLE_SIZE = 400
NUM_ITERS = 5
SEEDS = [101, 102, 103, 104, 105]


def run_experiment():
    print("\n" + "=" * 80)
    print("STARTING p-BART MINIMUM DRAW THRESHOLD & CONVERGENCE STUDY")
    print(f"Scenarios:   {TARGET_SCENARIOS}")
    print(f"Sample Size: N={SAMPLE_SIZE} | Replicates: {NUM_ITERS}")
    print(f"Configs:     {list(CONFIGURATIONS.keys())}")
    print(f"Output Dir:  {OUTPUT_DIR}")
    print("=" * 80 + "\n")
    
    raw_json_path = OUTPUT_DIR / "raw_results.json"
    existing_results = []
    if raw_json_path.exists():
        try:
            with open(raw_json_path, "r") as f:
                existing_results = json.load(f)
            for r in existing_results:
                if r.get("config_id") == "cfg3_draws1000_tune250":
                    r["config_id"] = "cfg3_draws1000_tune250_thin5"
                    r["config_label"] = "50 trees | 1,000 draws (thin 5 -> 800 draws) | 250 tune"
        except Exception as e:
            print(f"Notice: loading existing results: {e}")
            existing_results = []
            
    completed_map = {(r["scenario_id"], r["config_id"], r["seed"]): r for r in existing_results}
    all_results = []
    
    for scen_id in TARGET_SCENARIOS:
        scen_cfg = SCENARIOS[scen_id]
        print(f"\n==================== SCENARIO: {scen_cfg['name']} ({scen_id}) ====================")
        
        for rep_idx, seed in enumerate(SEEDS):
            print(f"\n--- [Replicate {rep_idx + 1}/{NUM_ITERS}] Seed {seed} ---")
            
            # Generate dataset
            train_data = generate_scenario_data(scen_cfg, n=SAMPLE_SIZE, seed=seed)
            test_data = None
            if "test_sample_size" in scen_cfg:
                test_data = generate_scenario_data(
                    scen_cfg,
                    n=scen_cfg["test_sample_size"],
                    seed=seed + 10000,
                    eval_times=train_data["eval_times"]
                )
            true_sv = train_data["sv_true_eval"]
            
            # 1. Run r-BART baseline for this replicate
            rb_key = (scen_id, "rbart_reference", seed)
            if rb_key in completed_map:
                print("  [r-BART Baseline] [Cached] Reusing previous result.")
                all_results.append(completed_map[rb_key])
            else:
                print("  [r-BART Baseline] Running...", end=" ", flush=True)
                rb_res = run_rbart(
                    train_data=train_data,
                    skip_r=False,
                    scenario_id=scen_id,
                    test_data=test_data,
                    causal_specs=scen_cfg.get("causal_contrasts", [])
                )
                rb_mets = compute_interval_metrics(
                    true_sv=true_sv,
                    pred_mean=rb_res["sv_mean"],
                    pred_ci_lower=rb_res["ci_lower"],
                    pred_ci_upper=rb_res["ci_upper"]
                )
                rb_diag = rb_res.get("diagnostics", {})
                rb_test_rmse = None
                if test_data and rb_res.get("test_results"):
                    rb_t_mets = compute_interval_metrics(
                        true_sv=test_data["sv_true_eval"],
                        pred_mean=rb_res["test_results"]["sv_mean"],
                        pred_ci_lower=rb_res["test_results"]["ci_lower"],
                        pred_ci_upper=rb_res["test_results"]["ci_upper"]
                    )
                    rb_test_rmse = rb_t_mets["overall_rmse"]
                    
                print(f"Done in {rb_res['timings']['total_seconds']:.2f}s | RMSE: {rb_mets['overall_rmse']:.4f} | R-hat: {rb_diag.get('rhat_mean', 0):.4f} | ESS: {rb_diag.get('ess_mean', 0):.1f}")
                
                all_results.append({
                    "scenario_id": scen_id,
                    "config_id": "rbart_reference",
                    "config_label": "r-BART (R::BART Baseline)",
                    "trees": 50,
                    "draws": 1000,
                    "tune": 250,
                    "chains": 4,
                    "seed": seed,
                    "model": "rbart",
                    "rmse": rb_mets["overall_rmse"],
                    "bias": rb_mets["overall_bias"],
                    "coverage": rb_mets.get("overall_coverage"),
                    "test_rmse": rb_test_rmse,
                    "rhat_mean": rb_diag.get("rhat_mean"),
                    "rhat_max": rb_diag.get("rhat_max"),
                    "ess_mean": rb_diag.get("ess_mean"),
                    "ess_min": rb_diag.get("ess_min"),
                    "fit_time": rb_res["timings"].get("fit_seconds"),
                    "total_time": rb_res["timings"]["total_seconds"],
                })
            # 2. Run p-BART configurations
            for cfg_id, m_cfg in CONFIGURATIONS.items():
                pb_key = (scen_id, cfg_id, seed)
                if pb_key in completed_map:
                    print(f"  [p-BART: {cfg_id}] [Cached] Reusing previous result.")
                    all_results.append(completed_map[pb_key])
                    continue
                    
                print(f"  [p-BART: {cfg_id}] Running...", end=" ", flush=True)
                pb_params = {k: v for k, v in m_cfg.items() if k != "label"}
                pb_res = run_pbart(
                    train_data=train_data,
                    model_params=pb_params,
                    causal_specs=scen_cfg.get("causal_contrasts", []),
                    test_data=test_data,
                    progressbar=False,
                    random_seed=seed
                )
                pb_mets = compute_interval_metrics(
                    true_sv=true_sv,
                    pred_mean=pb_res["sv_mean"],
                    pred_ci_lower=pb_res["ci_lower"],
                    pred_ci_upper=pb_res["ci_upper"]
                )
                pb_diag = pb_res.get("diagnostics", {})
                pb_test_rmse = None
                if test_data and pb_res.get("test_results"):
                    pb_t_mets = compute_interval_metrics(
                        true_sv=test_data["sv_true_eval"],
                        pred_mean=pb_res["test_results"]["sv_mean"],
                        pred_ci_lower=pb_res["test_results"]["ci_lower"],
                        pred_ci_upper=pb_res["test_results"]["ci_upper"]
                    )
                    pb_test_rmse = pb_t_mets["overall_rmse"]
                    
                print(f"Done in {pb_res['timings']['fit_seconds']:.2f}s | RMSE: {pb_mets['overall_rmse']:.4f} | R-hat: {pb_diag.get('rhat_mean', 0):.4f} | ESS: {pb_diag.get('ess_mean', 0):.1f}")
                
                all_results.append({
                    "scenario_id": scen_id,
                    "config_id": cfg_id,
                    "config_label": m_cfg["label"],
                    "trees": m_cfg["trees"],
                    "draws": m_cfg["draws"],
                    "tune": m_cfg["tune"],
                    "chains": m_cfg["chains"],
                    "seed": seed,
                    "model": "pbart",
                    "rmse": pb_mets["overall_rmse"],
                    "bias": pb_mets["overall_bias"],
                    "coverage": pb_mets.get("overall_coverage"),
                    "test_rmse": pb_test_rmse,
                    "rhat_mean": pb_diag.get("rhat_mean"),
                    "rhat_max": pb_diag.get("rhat_max"),
                    "ess_mean": pb_diag.get("ess_mean"),
                    "ess_min": pb_diag.get("ess_min"),
                    "fit_time": pb_res["timings"]["fit_seconds"],
                    "total_time": pb_res["timings"]["total_seconds"],
                })
                
    # Save raw json
    raw_json_path = OUTPUT_DIR / "raw_results.json"
    with open(raw_json_path, "w") as f:
        json.dump(all_results, f, indent=2)
    print(f"\nRaw results saved to: {raw_json_path}")
    
    # Generate Excel Report
    generate_excel_report(all_results)


def generate_excel_report(all_results: List[Dict[str, Any]]):
    df = pd.DataFrame(all_results)
    xlsx_path = OUTPUT_DIR / "p_bart_draw_threshold_analysis.xlsx"
    
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet
    
    # 1. Summary Matrix Tab
    ws_sum = wb.create_sheet(title="Threshold Summary Matrix")
    ws_sum.views.sheetView[0].showGridLines = True
    
    # Header styling
    navy_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    dark_gray = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
    accent_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    white_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    
    # Title
    ws_sum.merge_cells("A1:K1")
    t_cell = ws_sum.cell(1, 1, "p-BART MINIMUM DRAW THRESHOLD & CONVERGENCE STUDY")
    t_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    t_cell.fill = navy_fill
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 30
    
    ws_sum.merge_cells("A2:K2")
    sub_cell = ws_sum.cell(2, 1, "Systematic evaluation of MCMC draw/tune depth: Convergence (R-hat), Precision (ESS), Calibration (RMSE), and Runtime across 5 replicates (N=400).")
    sub_cell.font = Font(name="Calibri", size=10, italic=True)
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[2].height = 20
    
    headers = [
        "Scenario", "Configuration Preset", "Trees", "Draws / Chain", "Tune", "Fit Time (s)",
        "a-RMSE (Mean)", "95% Coverage", "R-hat (Mean)", "R-hat (Max)", "ESS Bulk (Mean)"
    ]
    
    for c_idx, h in enumerate(headers, 1):
        cell = ws_sum.cell(4, c_idx, h)
        cell.font = white_font
        cell.fill = dark_gray
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_sum.row_dimensions[4].height = 28
    
    curr_row = 5
    summary_data = []
    
    for scen_id in TARGET_SCENARIOS:
        scen_df = df[df["scenario_id"] == scen_id]
        
        # Add section header
        scen_title = SCENARIOS[scen_id]["name"]
        ws_sum.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=11)
        sec_cell = ws_sum.cell(curr_row, 1, f"SCENARIO: {scen_title.upper()}")
        sec_cell.font = bold_font
        sec_cell.fill = accent_fill
        sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws_sum.row_dimensions[curr_row].height = 22
        curr_row += 1
        
        config_order = ["rbart_reference"] + list(CONFIGURATIONS.keys())
        for cfg_id in config_order:
            sub = scen_df[scen_df["config_id"] == cfg_id]
            if sub.empty:
                continue
                
            cfg_label = sub["config_label"].iloc[0]
            trees = sub["trees"].iloc[0]
            draws = sub["draws"].iloc[0]
            tune = sub["tune"].iloc[0]
            
            fit_t = sub["fit_time"].mean()
            rmse_m = sub["rmse"].mean()
            cov_m = sub["coverage"].mean()
            rhat_m = sub["rhat_mean"].mean()
            rhat_max = sub["rhat_max"].mean()
            ess_m = sub["ess_mean"].mean()
            
            row_vals = [
                scen_id, cfg_label, trees, draws, tune,
                f"{fit_t:.2f}s" if fit_t else "N/A",
                f"{rmse_m:.4f}",
                f"{cov_m*100:.1f}%" if cov_m else "N/A",
                f"{rhat_m:.4f}" if rhat_m else "N/A",
                f"{rhat_max:.4f}" if rhat_max else "N/A",
                f"{ess_m:.0f}" if ess_m else "N/A"
            ]
            
            for c_idx, val in enumerate(row_vals, 1):
                cell = ws_sum.cell(curr_row, c_idx, val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center" if c_idx not in [1, 2] else "left", vertical="center")
                
                # Highlight convergence
                if c_idx in [9, 10] and rhat_m and float(rhat_m) > 1.05:
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid") # Red flag
                elif c_idx in [9, 10] and rhat_m and float(rhat_m) <= 1.01:
                    cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid") # Green good
                    
                if c_idx == 11 and ess_m and float(ess_m) >= 400:
                    cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                elif c_idx == 11 and ess_m and float(ess_m) < 200:
                    cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                    
            ws_sum.row_dimensions[curr_row].height = 20
            curr_row += 1
            
        curr_row += 1 # Spacing between scenarios
        
    for col in ws_sum.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_sum.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    # Save workbook
    wb.save(xlsx_path)
    print(f"Excel summary report created at: {xlsx_path}")
    
    # Copy to workspace root for user convenience
    root_copy = BASE_DIR.parent / "p_bart_draw_threshold_analysis.xlsx"
    wb.save(root_copy)
    print(f"Mirrored copy saved to workspace root: {root_copy}")


if __name__ == "__main__":
    run_experiment()
