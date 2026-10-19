#!/usr/bin/env python3
"""
Master Benchmark Excel Generator.
Consolidates active validation results across:
  - Simple 1_1 (1-Sample, 20% Censoring)
  - Simple 1_2 (1-Sample, 50% Censoring)
  - Simple 2_1 (2-Sample, 20% Censoring)
  - Simple 2_2 (2-Sample, 50% Censoring)
  - Complex 1 (Proportional Hazards, 10 Covariates)
  - Complex 2 (Non-Proportional Hazards, Crossing Hazards)
  - Complex 3 (Nonlinear U-Shape & Out-of-Sample Hold-Out Cohort)

Active Models (No legacy results):
  - p-BART (2026 Modernized)
  - r-BART (Live R::BART package execution)
  - CPHM / KM (Cox Proportional Hazards for Complex, Kaplan-Meier for Simple)

Structure:
  - Column 1: Metric Name
  - Column 2: Definition / Estimation Scope
  - Column 3: Analytical Truth
  - Columns Grouped by Output Type:
      - Means:   p-BART (2026), r-BART, CPHM / KM
      - Medians: p-BART (2026), r-BART, CPHM / KM
      - SDs:     p-BART (2026), r-BART, CPHM / KM
"""

import sys
import argparse
import json
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent.parent
RESULTS_DIR = BASE_DIR / "results"
OUTPUT_XLSX = RESULTS_DIR / "master_benchmark_results.xlsx"


def load_json(p: Path):
    if not p.exists():
        return None
    with open(p, "r") as f:
        return json.load(f)


def safe_stats(arr):
    """Returns (mean, median, sd) for an array or list."""
    if arr is None or len(arr) == 0:
        return (None, None, None)
    a = np.asarray(arr, dtype=np.float64)
    a = a[~np.isnan(a)]
    if len(a) == 0:
        return (None, None, None)
    return (float(np.mean(a)), float(np.median(a)), float(np.std(a)))


def get_scenario_data(experiment_name, scenario_id):
    """Loads replicate records for a scenario from results."""
    suite_dir = RESULTS_DIR / experiment_name
    for scen_dir in suite_dir.glob(f"{scenario_id}_*"):
        rep_file = scen_dir / "replicate_metrics.json"
        causal_file = scen_dir / "causal_effects.json"
        if rep_file.exists():
            reps = load_json(rep_file)
            causal = load_json(causal_file) if causal_file.exists() else None
            return reps, causal
    return None, None


def style_worksheet(ws):
    """Applies clean, professional corporate typography and cell styling."""
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid") # Deep navy
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    subgroup_fill = PatternFill(start_color="E8EEF5", end_color="E8EEF5", fill_type="solid") # Soft ice blue
    subgroup_font = Font(name="Calibri", size=11, bold=True, color="1B365D")
    
    section_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    section_font = Font(name="Calibri", size=10, bold=True, color="333333")
    
    regular_font = Font(name="Calibri", size=10, color="000000")
    desc_font = Font(name="Calibri", size=9, italic=True, color="555555")
    truth_font = Font(name="Calibri", size=10, bold=True, color="004D40") # Dark Teal
    
    thin_border_side = Side(border_style="thin", color="D0D5DD")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(bottom=Side(border_style="medium", color="1B365D"), left=thin_border_side, right=thin_border_side, top=thin_border_side)
    
    # Header row formatting (Row 1: Super-headers, Row 2: Sub-headers)
    for col in range(1, ws.max_column + 1):
        cell1 = ws.cell(row=1, column=col)
        cell2 = ws.cell(row=2, column=col)
        cell1.fill = header_fill
        cell1.font = header_font
        cell1.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell2.fill = subgroup_fill
        cell2.font = subgroup_font
        cell2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell1.border = thin_border
        cell2.border = thick_bottom

    # Row formatting
    for row in range(3, ws.max_row + 1):
        metric_name = ws.cell(row=row, column=1).value
        is_section_header = metric_name and str(metric_name).startswith("---")
        
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            
            if is_section_header:
                cell.fill = section_fill
                cell.font = section_font
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                if col == 1:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="111111")
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif col == 2:
                    cell.font = desc_font
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                elif col == 3:
                    cell.font = truth_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = regular_font
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    # Format numbers nicely
                    val = cell.value
                    if isinstance(val, (int, float)):
                        if abs(val) < 0.0001 and val != 0:
                            cell.number_format = "0.00000"
                        elif "Coverage" in str(metric_name) or "%" in str(metric_name):
                            cell.number_format = "0.00%" if val <= 1.0 else "0.00"
                        elif "Time" in str(metric_name) or "Seconds" in str(metric_name):
                            cell.number_format = "0.00"
                        else:
                            cell.number_format = "0.0000"

    # Set column widths
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 16
    for c in ["D", "E", "F", "G", "H", "I", "J", "K", "L"]:
        ws.column_dimensions[c].width = 18
        
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 24


def build_scenario_tab(wb, title, tab_name, reps, causal, category, is_km=False):
    """Builds a formatted scenario tab with active models only."""
    ws = wb.create_sheet(title=tab_name)
    
    other_label = "KM" if is_km else "CoxPH"
    
    # 2-Row Grouped Header Structure
    row1 = [
        "METRIC INFORMATION", "", "",
        "MEAN (ACROSS REPLICATES)", "", "",
        "MEDIAN (ACROSS REPLICATES)", "", "",
        "STANDARD DEVIATION (ACROSS REPLICATES)", "", ""
    ]
    row2 = [
        "Metric Name", "Definition / Estimation Scope", "Analytical Truth",
        "p-BART (2026)", "r-BART", other_label,
        "p-BART (2026)", "r-BART", other_label,
        "p-BART (2026)", "r-BART", other_label,
    ]
    
    ws.append(row1)
    ws.append(row2)
    
    # Merge group headers
    ws.merge_cells("A1:C1")
    ws.merge_cells("D1:F1")
    ws.merge_cells("G1:I1")
    ws.merge_cells("J1:L1")
    
    # Separate replicates by model
    pb_reps = [r for r in reps if r.get("model") == "pbart"] if reps else []
    rb_reps = [r for r in reps if r.get("model") == "rbart"] if reps else []
    oth_reps = [r for r in reps if r.get("model") in ("km", "coxph")] if reps else []
    
    def add_row(metric_name, desc, truth, s_pb, s_rb, s_oth):
        """Adds a data row with Means, Medians, and SDs aligned."""
        mean_pb, med_pb, sd_pb = s_pb
        mean_rb, med_rb, sd_rb = s_rb
        mean_oth, med_oth, sd_oth = s_oth
        
        ws.append([
            metric_name, desc, truth,
            mean_pb, mean_rb, mean_oth,
            med_pb, med_rb, med_oth,
            sd_pb, sd_rb, sd_oth
        ])
        
    def add_header(section_title):
        ws.append([f"--- {section_title.upper()} ---", "", "", "", "", "", "", "", "", "", "", ""])

    # 1. OVERALL CALIBRATION
    add_header("Overall Survival Curve Calibration Across All Evaluation Timepoints")
    add_row(
        "Overall Survival RMSE (a-RMSE)",
        "Root Mean Squared Error comparing predicted survival probability S_hat(t|X) to true analytical survival S(t|X) pooled across all subjects and time intervals.",
        0.0000,
        safe_stats([r.get("rmse") for r in pb_reps]),
        safe_stats([r.get("rmse") for r in rb_reps]),
        safe_stats([r.get("rmse") for r in oth_reps])
    )
    add_row(
        "Overall Signed Calibration Bias",
        "Average pointwise difference [S_hat(t|X) - S(t|X)]. Positive indicates overestimating survival probability; negative indicates underestimating survival.",
        0.0000,
        safe_stats([r.get("bias") for r in pb_reps]),
        safe_stats([r.get("bias") for r in rb_reps]),
        safe_stats([r.get("bias") for r in oth_reps])
    )
    add_row(
        "Empirical 95% Credible / Confidence Coverage",
        "Proportion of true survival probabilities S(t|X) contained inside the estimated 95% credible interval [L_95, U_95]. Targeted nominal coverage is exactly 95.00%.",
        0.9500,
        safe_stats([r.get("coverage") for r in pb_reps]),
        safe_stats([r.get("coverage") for r in rb_reps]),
        safe_stats([r.get("coverage") for r in oth_reps])
    )
    add_row(
        "95% Credible Interval Width (Precision)",
        "Mean length of the 95% credible/confidence interval (U_95 - L_95). Narrower width indicates higher precision / lower posterior uncertainty.",
        "Narrower",
        safe_stats([r.get("interval_length") for r in pb_reps]),
        safe_stats([r.get("interval_length") for r in rb_reps]),
        safe_stats([r.get("interval_length") for r in oth_reps])
    )

    # 2. TIME-SPECIFIC INTERVAL METRICS
    pcts = [10, 25, 50, 75, 90]
    add_header("Pointwise Calibration Error (a-RMSE) Across Survival Follow-up Percentiles")
    for p_idx, p_val in enumerate(pcts):
        # Extract per-interval arrays
        pb_inv = [r["rmse_per_interval"][p_idx + 1] if "complex" in category and len(r.get("rmse_per_interval", [])) > p_idx + 1 else (r["rmse_per_interval"][p_idx] if len(r.get("rmse_per_interval", [])) > p_idx else None) for r in pb_reps if r.get("rmse_per_interval")]
        rb_inv = [r["rmse_per_interval"][p_idx + 1] if "complex" in category and len(r.get("rmse_per_interval", [])) > p_idx + 1 else (r["rmse_per_interval"][p_idx] if len(r.get("rmse_per_interval", [])) > p_idx else None) for r in rb_reps if r.get("rmse_per_interval")]
        oth_inv = [r["rmse_per_interval"][p_idx + 1] if "complex" in category and len(r.get("rmse_per_interval", [])) > p_idx + 1 else (r["rmse_per_interval"][p_idx] if len(r.get("rmse_per_interval", [])) > p_idx else None) for r in oth_reps if r.get("rmse_per_interval")]
        
        add_row(
            f"RMSE at {p_val}% Survival Time",
            f"Pointwise root mean squared error evaluating calibrated survival at the {p_val}th percentile of the population survival curve.",
            0.0000,
            safe_stats(pb_inv),
            safe_stats(rb_inv),
            safe_stats(oth_inv)
        )

    add_header("Pointwise Signed Calibration Bias Across Survival Follow-up Percentiles")
    for p_idx, p_val in enumerate(pcts):
        pb_b_inv = [r["bias_per_interval"][p_idx + 1] if "complex" in category and len(r.get("bias_per_interval", [])) > p_idx + 1 else (r["bias_per_interval"][p_idx] if len(r.get("bias_per_interval", [])) > p_idx else None) for r in pb_reps if r.get("bias_per_interval")]
        rb_b_inv = [r["bias_per_interval"][p_idx + 1] if "complex" in category and len(r.get("bias_per_interval", [])) > p_idx + 1 else (r["bias_per_interval"][p_idx] if len(r.get("bias_per_interval", [])) > p_idx else None) for r in rb_reps if r.get("bias_per_interval")]
        oth_b_inv = [r["bias_per_interval"][p_idx + 1] if "complex" in category and len(r.get("bias_per_interval", [])) > p_idx + 1 else (r["bias_per_interval"][p_idx] if len(r.get("bias_per_interval", [])) > p_idx else None) for r in oth_reps if r.get("bias_per_interval")]
        
        add_row(
            f"Signed Bias at {p_val}% Survival Time",
            f"Signed bias [S_hat(t) - S(t)] at the {p_val}th percentile time. Highlights early vs late hazard profile calibration.",
            0.0000,
            safe_stats(pb_b_inv),
            safe_stats(rb_b_inv),
            safe_stats(oth_b_inv)
        )

    # 3. RISK-STRATIFIED SUBGROUP CALIBRATION (TOP 25% vs BOTTOM 25%)
    has_strat_reps = any("rmse_top25" in r for r in pb_reps or rb_reps)
    if has_strat_reps:
        add_header("Risk-Stratified Subgroup Calibration (Top 25% Low-Risk & Bottom 25% High-Risk)")
        add_row(
            "Low-Risk Subgroup (Top 25% Survival) RMSE",
            "Root Mean Squared Error of the marginal survival curve for the top 25% highest-survival subjects.",
            0.0000,
            safe_stats([r.get("rmse_top25") for r in pb_reps]),
            safe_stats([r.get("rmse_top25") for r in rb_reps]),
            safe_stats([r.get("rmse_top25") for r in oth_reps])
        )
        add_row(
            "Low-Risk Subgroup (Top 25% Survival) Signed Bias",
            "Signed calibration bias for the top 25% highest-survival subjects. Detects floor/ceiling effects.",
            0.0000,
            safe_stats([r.get("bias_top25") for r in pb_reps]),
            safe_stats([r.get("bias_top25") for r in rb_reps]),
            safe_stats([r.get("bias_top25") for r in oth_reps])
        )
        add_row(
            "Low-Risk Subgroup (Top 25% Survival) 95% Coverage",
            "Proportion of time true top-25% marginal survival curve is enclosed within the 95% credible interval across intervals.",
            0.9500,
            safe_stats([r.get("coverage_top25") for r in pb_reps]),
            safe_stats([r.get("coverage_top25") for r in rb_reps]),
            (None, None, None)
        )
        add_row(
            "High-Risk Subgroup (Bottom 25% Survival) RMSE",
            "Root Mean Squared Error of the marginal survival curve for the bottom 25% lowest-survival subjects.",
            0.0000,
            safe_stats([r.get("rmse_bot25") for r in pb_reps]),
            safe_stats([r.get("rmse_bot25") for r in rb_reps]),
            safe_stats([r.get("rmse_bot25") for r in oth_reps])
        )
        add_row(
            "High-Risk Subgroup (Bottom 25% Survival) Signed Bias",
            "Signed calibration bias for the bottom 25% lowest-survival subjects. Detects premature hazard decay.",
            0.0000,
            safe_stats([r.get("bias_bot25") for r in pb_reps]),
            safe_stats([r.get("bias_bot25") for r in rb_reps]),
            safe_stats([r.get("bias_bot25") for r in oth_reps])
        )
        add_row(
            "High-Risk Subgroup (Bottom 25% Survival) 95% Coverage",
            "Proportion of time true bottom-25% marginal survival curve is enclosed within the 95% credible interval across intervals.",
            0.9500,
            safe_stats([r.get("coverage_bot25") for r in pb_reps]),
            safe_stats([r.get("coverage_bot25") for r in rb_reps]),
            (None, None, None)
        )
        add_row(
            "Extreme High-Risk Boundary (Min Subject) 95% Coverage",
            "Pointwise 95% credible interval coverage for the single individual with lowest predicted/true survival.",
            0.9500,
            safe_stats([r.get("coverage_min") for r in pb_reps]),
            safe_stats([r.get("coverage_min") for r in rb_reps]),
            (None, None, None)
        )
        add_row(
            "Extreme Low-Risk Boundary (Max Subject) 95% Coverage",
            "Pointwise 95% credible interval coverage for the single individual with highest predicted/true survival.",
            0.9500,
            safe_stats([r.get("coverage_max") for r in pb_reps]),
            safe_stats([r.get("coverage_max") for r in rb_reps]),
            (None, None, None)
        )

    # 4. OUT-OF-SAMPLE TEST COHORT (For Complex 3)
    has_test_cohort = any("test_rmse" in r for r in pb_reps or rb_reps)
    if has_test_cohort:
        add_header("Out-of-Sample Hold-Out Test Cohort Generalization (Unseen Subjects)")
        add_row(
            "Hold-out Test RMSE (Overall)",
            "Out-of-sample calibration error evaluated on a completely independent test cohort generated from the identical non-linear DGP.",
            0.0000,
            safe_stats([r.get("test_rmse") for r in pb_reps]),
            safe_stats([r.get("test_rmse") for r in rb_reps]),
            safe_stats([r.get("test_rmse") for r in oth_reps])
        )
        add_row(
            "Hold-out Test 95% Credible Coverage",
            "Empirical 95% coverage on unseen hold-out subjects across follow-up time deciles. Reflects true epistemic uncertainty calibration.",
            0.9500,
            safe_stats([r.get("test_coverage") for r in pb_reps]),
            safe_stats([r.get("test_coverage") for r in rb_reps]),
            safe_stats([r.get("test_coverage") for r in oth_reps])
        )
        
        # Pointwise deciles
        dec_indices = {0.1: 1, 0.3: 3, 0.5: 5, 0.7: 7, 0.9: 9}
        for d in [0.1, 0.3, 0.5, 0.7, 0.9]:
            col_idx = dec_indices[d]
            pb_vals_rmse = [r["test_rmse_per_interval"][col_idx] for r in pb_reps if "test_rmse_per_interval" in r and len(r["test_rmse_per_interval"]) > col_idx]
            rb_vals_rmse = [r["test_rmse_per_interval"][col_idx] for r in rb_reps if "test_rmse_per_interval" in r and len(r["test_rmse_per_interval"]) > col_idx]
            oth_vals_rmse = [r["test_rmse_per_interval"][col_idx] for r in oth_reps if "test_rmse_per_interval" in r and len(r["test_rmse_per_interval"]) > col_idx]
            
            pb_vals_bias = [r["test_bias_per_interval"][col_idx] for r in pb_reps if "test_bias_per_interval" in r and len(r["test_bias_per_interval"]) > col_idx]
            rb_vals_bias = [r["test_bias_per_interval"][col_idx] for r in rb_reps if "test_bias_per_interval" in r and len(r["test_bias_per_interval"]) > col_idx]
            oth_vals_bias = [r["test_bias_per_interval"][col_idx] for r in oth_reps if "test_bias_per_interval" in r and len(r["test_bias_per_interval"]) > col_idx]
            
            add_row(
                f"Test RMSE at Time Decile {d}",
                f"Out-of-sample error at follow-up time decile {d}. Highlights tail anchor performance (deciles 0.1 & 0.9).",
                0.0000,
                safe_stats(pb_vals_rmse),
                safe_stats(rb_vals_rmse),
                safe_stats(oth_vals_rmse)
            )
            add_row(
                f"Test Signed Bias at Time Decile {d}",
                f"Out-of-sample signed bias at follow-up time decile {d}.",
                0.0000,
                safe_stats(pb_vals_bias),
                safe_stats(rb_vals_bias),
                safe_stats(oth_vals_bias)
            )

    # 5. CAUSAL ESTIMANDS
    if causal:
        add_header("Causal Effect Estimands (Hazard Ratios & Risk Differences)")
        # Identify variables in causal records
        unique_vars = sorted(list({r.get("variable") for r in causal if r.get("variable")}))
        for var_name in unique_vars:
            sub_c = [r for r in causal if r.get("variable") == var_name]
            truth_val = next((r.get("true_hr") for r in sub_c if r.get("true_hr") is not None), None)
            
            pb_causal = [r["hazard_ratio_weighted_mean"] for r in sub_c if r.get("model") == "pbart" and r.get("hazard_ratio_weighted_mean") is not None]
            rb_causal = [r["hazard_ratio_weighted_mean"] for r in sub_c if r.get("model") == "rbart" and r.get("hazard_ratio_weighted_mean") is not None]
            oth_causal = [r["hazard_ratio_weighted_mean"] for r in sub_c if r.get("model") == "coxph" and r.get("hazard_ratio_weighted_mean") is not None]
            
            add_row(
                f"{var_name} Marginal Causal Hazard Ratio",
                f"G-computation marginal causal hazard ratio across the population cohort for treatment indicator {var_name}.",
                truth_val if truth_val is not None else "N/A",
                safe_stats(pb_causal),
                safe_stats(rb_causal),
                safe_stats(oth_causal)
            )
            
            pb_hr_w = [r["hazard_ratio_ci_width"] for r in sub_c if r.get("model") == "pbart" and r.get("hazard_ratio_ci_width") is not None]
            rb_hr_w = [r["hazard_ratio_ci_width"] for r in sub_c if r.get("model") == "rbart" and r.get("hazard_ratio_ci_width") is not None]
            oth_hr_w = [r["hazard_ratio_ci_width"] for r in sub_c if r.get("model") == "coxph" and r.get("hazard_ratio_ci_width") is not None]
            add_row(
                f"{var_name} Marginal Hazard Ratio 95% Interval Width (Precision)",
                f"Average width of the 95% credible/confidence interval (U_95 - L_95) for the marginal hazard ratio.",
                "Narrower",
                safe_stats(pb_hr_w),
                safe_stats(rb_hr_w),
                safe_stats(oth_hr_w)
            )
            
            pb_hr_cov = [r["coverage_hr"] for r in sub_c if r.get("model") == "pbart" and r.get("coverage_hr") is not None]
            rb_hr_cov = [r["coverage_hr"] for r in sub_c if r.get("model") == "rbart" and r.get("coverage_hr") is not None]
            oth_hr_cov = [r["coverage_hr"] for r in sub_c if r.get("model") == "coxph" and r.get("coverage_hr") is not None]
            add_row(
                f"{var_name} Marginal HR Empirical 95% CI Coverage",
                f"Empirical frequency of the 95% credible/confidence interval enclosing the analytical true marginal Hazard Ratio across iterations.",
                0.9500,
                safe_stats(pb_hr_cov),
                safe_stats(rb_hr_cov),
                safe_stats(oth_hr_cov)
            )
            
            true_rd_val = next((r.get("true_rd_at_t50") for r in sub_c if r.get("true_rd_at_t50") is not None), None)
            pb_rd = [r["risk_difference_t50_mean"] for r in sub_c if r.get("model") == "pbart" and r.get("risk_difference_t50_mean") is not None]
            rb_rd = [r["risk_difference_t50_mean"] for r in sub_c if r.get("model") == "rbart" and r.get("risk_difference_t50_mean") is not None]
            oth_rd = [r["risk_difference_t50_mean"] for r in sub_c if r.get("model") == "coxph" and r.get("risk_difference_t50_mean") is not None]
            add_row(
                f"{var_name} Marginal Risk Difference at Median Horizon RD(t50)",
                f"Counterfactual excess event risk RD(t50) = S0(t50) - S1(t50) evaluated at median follow-up time.",
                true_rd_val if true_rd_val is not None else "N/A",
                safe_stats(pb_rd),
                safe_stats(rb_rd),
                safe_stats(oth_rd)
            )
            
            pb_rd_w = [r["risk_difference_t50_ci_width"] for r in sub_c if r.get("model") == "pbart" and r.get("risk_difference_t50_ci_width") is not None]
            rb_rd_w = [r["risk_difference_t50_ci_width"] for r in sub_c if r.get("model") == "rbart" and r.get("risk_difference_t50_ci_width") is not None]
            oth_rd_w = [r["risk_difference_t50_ci_width"] for r in sub_c if r.get("model") == "coxph" and r.get("risk_difference_t50_ci_width") is not None]
            add_row(
                f"{var_name} Marginal RD(t50) 95% Interval Width (Precision)",
                f"Average width of the 95% credible/confidence interval (U_95 - L_95) for the marginal risk difference at median follow-up time.",
                "Narrower",
                safe_stats(pb_rd_w),
                safe_stats(rb_rd_w),
                safe_stats(oth_rd_w)
            )
            
            pb_rd_cov = [r["coverage_rd_t50"] for r in sub_c if r.get("model") == "pbart" and r.get("coverage_rd_t50") is not None]
            rb_rd_cov = [r["coverage_rd_t50"] for r in sub_c if r.get("model") == "rbart" and r.get("coverage_rd_t50") is not None]
            oth_rd_cov = [r["coverage_rd_t50"] for r in sub_c if r.get("model") == "coxph" and r.get("coverage_rd_t50") is not None]
            add_row(
                f"{var_name} Marginal RD(t50) Empirical 95% CI Coverage",
                f"Empirical coverage frequency of the 95% confidence/credible interval enclosing true RD(t50) across simulation iterations.",
                0.9500,
                safe_stats(pb_rd_cov),
                safe_stats(rb_rd_cov),
                safe_stats(oth_rd_cov)
            )

    # 5. MCMC CONVERGENCE & EFFECTIVE SAMPLING DIAGNOSTICS (R-HAT & ESS)
    has_diag = any("rhat_mean" in r and r["rhat_mean"] is not None for r in pb_reps + rb_reps)
    if has_diag:
        add_header("MCMC Convergence & Effective Sampling Diagnostics (R-hat & ESS)")
        add_row(
            "Gelman-Rubin Convergence Diagnostic (R-hat Mean)",
            "Potential scale reduction factor R-hat across parallel MCMC chains evaluating cohort survival. Ideal value is 1.00; values < 1.05 indicate robust inter-chain convergence.",
            1.0000,
            safe_stats([r.get("rhat_mean") for r in pb_reps]),
            safe_stats([r.get("rhat_mean") for r in rb_reps]),
            (None, None, None)
        )
        add_row(
            "Worst-Case Convergence Diagnostic (R-hat Max)",
            "Maximum R-hat observed across all discrete survival time intervals. Detects localized non-convergence or chain stickiness in the tails.",
            1.0000,
            safe_stats([r.get("rhat_max") for r in pb_reps]),
            safe_stats([r.get("rhat_max") for r in rb_reps]),
            (None, None, None)
        )
        add_row(
            "Effective Sample Size Bulk (ESS Mean)",
            "Estimated number of independent posterior samples accounting for autocorrelation across chains (mean across time intervals). Threshold > 400 recommended for stable posterior inference.",
            "Higher (>400)",
            safe_stats([r.get("ess_bulk") for r in pb_reps]),
            safe_stats([r.get("ess_bulk") for r in rb_reps]),
            (None, None, None)
        )
        add_row(
            "Effective Sample Size Worst-Case (ESS Min)",
            "Minimum effective sample size across all discrete intervals. Identifies lower-bound sampling precision in sparse tail intervals.",
            "Higher (>100)",
            safe_stats([r.get("ess_min") for r in pb_reps]),
            safe_stats([r.get("ess_min") for r in rb_reps]),
            (None, None, None)
        )

    # 6. COMPUTATIONAL PERFORMANCE
    add_header("Computational Performance & Sampling Runtime")
    add_row(
        "MCMC Model Fit Time (Seconds)",
        "Raw wall-clock time required to compile the model, initialize trees, and complete posterior MCMC sampling across all chains.",
        "Lower",
        safe_stats([r.get("fit_time_seconds") for r in pb_reps]),
        safe_stats([r.get("fit_time_seconds") for r in rb_reps]),
        safe_stats([r.get("fit_time_seconds") for r in oth_reps])
    )
    add_row(
        "Total Pipeline Execution Time (Seconds)",
        "End-to-end wall-clock time including data preprocessing, long-time conversion, MCMC sampling, and posterior survival curve prediction.",
        "Lower",
        safe_stats([r.get("total_time_seconds") for r in pb_reps]),
        safe_stats([r.get("total_time_seconds") for r in rb_reps]),
        safe_stats([r.get("total_time_seconds") for r in oth_reps])
    )

    style_worksheet(ws)


def build_overview_tab(wb, experiment_name, scenarios):
    """Creates an Executive Summary tab."""
    ws = wb.create_sheet(title="Executive Summary", index=0)
    
    # Title Block
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = "BART SURVIVAL VALIDATION SUITE: MASTER BENCHMARK RESULTS"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # Narrative Summary
    ws.merge_cells("A3:I3")
    ws["A3"] = f"This master report consolidates active validation benchmark results comparing modernized 2026 p-BART (SurvBART) against live r-BART (R::BART package) and CoxPH/KM baselines (Suite: {experiment_name})."
    ws["A3"].font = Font(name="Calibri", size=10, italic=True)

    headers = [
        "Scenario Name", "DGP Characteristics", "Sample Size (N)", "Follow-up Horizon",
        "p-BART (2026) RMSE", "r-BART RMSE", "Cox / KM RMSE",
        "p-BART Fit Time (s)", "r-BART Fit Time (s)"
    ]
    ws.append([])
    ws.append(headers)
    ws.row_dimensions[5].height = 25
    
    # Format table header
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    for col in range(1, 10):
        cell = ws.cell(row=5, column=col)
        cell.fill = header_fill
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_idx = 6
    for tab_name, scen_id, full_title, category, pref_key, horizon in scenarios:
        reps, _ = get_scenario_data(experiment_name, scen_id)
        pb_reps = [r for r in reps if r.get("model") == "pbart"] if reps else []
        rb_reps = [r for r in reps if r.get("model") == "rbart"] if reps else []
        oth_reps = [r for r in reps if r.get("model") in ("km", "coxph")] if reps else []
        
        m_pb_rmse, _, _ = safe_stats([r.get("rmse") for r in pb_reps])
        m_rb_rmse, _, _ = safe_stats([r.get("rmse") for r in rb_reps])
        m_oth_rmse, _, _ = safe_stats([r.get("rmse") for r in oth_reps])
        
        m_pb_time, _, _ = safe_stats([r.get("fit_time_seconds") for r in pb_reps])
        m_rb_time, _, _ = safe_stats([r.get("fit_time_seconds") for r in rb_reps])
        
        ws.append([
            full_title,
            category.replace("_", " ").title(),
            "400",
            horizon,
            f"{m_pb_rmse:.4f}" if m_pb_rmse is not None else "N/A",
            f"{m_rb_rmse:.4f}" if m_rb_rmse is not None else "N/A",
            f"{m_oth_rmse:.4f}" if m_oth_rmse is not None else "N/A",
            f"{m_pb_time:.2f}s" if m_pb_time is not None else "N/A",
            f"{m_rb_time:.2f}s" if m_rb_time is not None else "N/A"
        ])
        row_idx += 1

    # Style Executive Summary
    thin_border_side = Side(border_style="thin", color="D0D5DD")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    for r in range(6, row_idx):
        for c in range(1, 10):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            if c in (1, 2):
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 20
    ws.column_dimensions["H"].width = 20
    ws.column_dimensions["I"].width = 20


def build_metric_dictionary(wb):
    """Adds a comprehensive Metric Dictionary tab defining every metric, estimand, and target."""
    ws = wb.create_sheet(title="Metric Dictionary")
    
    headers = ["Metric Term", "Target Value / Ideal", "Mathematical Formulation", "Clinical / Methodological Interpretation"]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    for col in range(1, 5):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    dictionary_data = [
        (
            "Overall Survival RMSE (a-RMSE)",
            "0.0000 (Minimization)",
            "sqrt( 1/(N*K) * sum_{i,k} [S_hat(t_k|X_i) - S(t_k|X_i)]^2 )",
            "Root mean squared error measuring global calibration accuracy across all individuals and follow-up time intervals."
        ),
        (
            "Overall Signed Calibration Bias",
            "0.0000 (Unbiased)",
            "1/(N*K) * sum_{i,k} [S_hat(t_k|X_i) - S(t_k|X_i)]",
            "Signed average divergence. Positive bias indicates overestimating true survival probability; negative indicates underestimation."
        ),
        (
            "Empirical 95% Credible / Confidence Coverage",
            "0.9500 (95.0%)",
            "1/(N*K) * sum_{i,k} 1{ L_{0.025}(t_k|X_i) <= S(t_k|X_i) <= U_{0.975}(t_k|X_i) }",
            "Evaluates epistemic uncertainty quantification. Reflects whether the 95% posterior credible interval encloses true survival 95% of the time."
        ),
        (
            "95% Credible Interval Width (Precision)",
            "Narrower is superior at nominal coverage",
            "1/(N*K) * sum_{i,k} [U_{0.975}(t_k|X_i) - L_{0.025}(t_k|X_i)]",
            "Mean length of the posterior uncertainty band. Measures statistical efficiency and epistemic certainty."
        ),
        (
            "Pointwise Percentile RMSE & Bias (10% to 90%)",
            "0.0000 across all percentiles",
            "sqrt( 1/N * sum_i [S_hat(t_pct|X_i) - S(t_pct|X_i)]^2 )",
            "Evaluates calibration at specific survival milestones (early, median, and late survival). Detects time-dependent miscalibration."
        ),
        (
            "Hold-out Test RMSE & Coverage",
            "0.0000 RMSE / 0.9500 Coverage",
            "Evaluated on independent unseen test cohort N=400",
            "Measures generalization to unseen patients, guarding against tree overfitting on non-linear response surfaces."
        ),
        (
            "Marginal Causal Hazard Ratio (G-Computation)",
            "Exact Analytical True HR",
            "exp( E_X [ log h(t|X, A=1) - log h(t|X, A=0) ] )",
            "Population average causal contrast for counterfactual treatment assignment. Estimated via posterior predictive counterfactual integration."
        ),
        (
            "MCMC Model Fit Time (Seconds)",
            "Lower is superior",
            "Wall-clock fit duration",
            "Computation time for MCMC tree posterior sampling."
        ),
        (
            "Total Pipeline Execution Time (Seconds)",
            "Lower is superior",
            "End-to-end execution duration",
            "Total run time including long data transformation, model fitting, and survival prediction."
        )
    ]

    for row in dictionary_data:
        ws.append(list(row))

    thin_border_side = Side(border_style="thin", color="D0D5DD")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    for r in range(2, len(dictionary_data) + 2):
        ws.row_dimensions[r].height = 28
        for c in range(1, 5):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border
            if c == 1:
                cell.font = Font(name="Calibri", size=10, bold=True, color="111111")
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c == 2:
                cell.font = Font(name="Calibri", size=10, bold=True, color="004D40")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c == 3:
                cell.font = Font(name="Courier New", size=9, color="333333")
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.font = Font(name="Calibri", size=9.5, italic=True, color="444444")
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 58


def main():
    parser = argparse.ArgumentParser(description="Generate Master Excel Report")
    parser.add_argument("--experiment-name", type=str, default="suite_400_10iter_live_rbart", help="Benchmark experiment folder name")
    parser.add_argument("--output-filename", type=str, default=None, help="Custom output filename saved in results")
    args = parser.parse_args()
    
    experiment_name = args.experiment_name
    print(f"\n" + "=" * 75)
    print(f"GENERATING MASTER BENCHMARK EXCEL REPORT: {experiment_name}")
    print("=" * 75)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    scenarios = [
        # (tab_name, scenario_id, title, category, prefix_or_key, horizon)
        ("Simple 1_1 (20% Cens)", "simple_1_1", "Simple 1-Sample (20% Censoring)", "simple", "1_1", "180 days"),
        ("Simple 1_2 (50% Cens)", "simple_1_2", "Simple 1-Sample (50% Censoring)", "simple", "1_2", "180 days"),
        ("Simple 2_1 (20% Cens)", "simple_2_1", "Simple 2-Sample (20% Censoring)", "simple", "2_1", "180 days"),
        ("Simple 2_2 (50% Cens)", "simple_2_2", "Simple 2-Sample (50% Censoring)", "simple", "2_2", "180 days"),
        ("Complex 1 (PH Linear)", "complex_1", "Complex 1 (Proportional Hazards)", "complex_binary", "cmplx1", "5.0 years"),
        ("Complex 2 (Non-PH)", "complex_2", "Complex 2 (Non-Proportional Hazards)", "complex_binary", "cmplx2", "5.0 years"),
        ("Complex 3 (Nonlinear)", "complex_3", "Complex 3 (Nonlinear Generalization)", "complex_continuous", "cmplx3", "5.0 years"),
    ]

    for tab_name, scen_id, full_title, category, pref_key, horizon in scenarios:
        print(f"  -> Building tab: {tab_name}...")
        reps, causal = get_scenario_data(experiment_name, scen_id)
        is_km = (category == "simple")
        build_scenario_tab(wb, full_title, tab_name, reps, causal, category, is_km=is_km)

    # 3. Add Overview & Metric Dictionary Tabs
    print("  -> Building Overview & Executive Summary tab...")
    build_overview_tab(wb, experiment_name, scenarios)
    
    print("  -> Building Metric Dictionary tab...")
    build_metric_dictionary(wb)

    # 4. Save workbook
    if args.output_filename:
        out_path = RESULTS_DIR / args.output_filename
    else:
        out_path = OUTPUT_XLSX
        
    wb.save(out_path)
    print(f"\nSUCCESS: Master Excel report created at: {out_path}")
    
    # Also mirror copy to workspace root if standard master
    if not args.output_filename:
        root_copy = BASE_DIR.parent / "BART_Survival_Master_Benchmark_Results.xlsx"
        wb.save(root_copy)
        print(f"Mirrored copy saved to workspace root: {root_copy}")
    print("=" * 75 + "\n")


if __name__ == "__main__":
    main()
